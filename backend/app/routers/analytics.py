"""
Analytics Router
Handles GET endpoints for USER, PREMIUM_USER, and ADMIN analytics.
Role enforcement is done at both the service layer and dependency level.
"""
import io
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.user import User
from app.routers.auth import get_current_user, require_premium_or_admin
from app.services import analytics_service


router = APIRouter(
    prefix="/analytics",
    tags=["Analytics"]
)


# ==========================================
# SPENDING BY CATEGORY
# (All roles — date range ignored for USER)
# ==========================================

@router.get("/spending-by-category")
def spending_by_category(
    start_date: Optional[date] = Query(None, description="YYYY-MM-DD (Premium only)"),
    end_date: Optional[date] = Query(None, description="YYYY-MM-DD (Premium only)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Returns spending by category.
    - USER: always current month, date params ignored.
    - PREMIUM/ADMIN: uses provided date range or defaults to current month.
    """
    data = analytics_service.get_spending_by_category(
        db=db,
        user_id=current_user.user_id,
        start_date=start_date,
        end_date=end_date,
        role=current_user.role
    )
    return data


# ==========================================
# SUMMARY (Income / Expense / Net Balance)
# (All roles — date range ignored for USER)
# ==========================================

@router.get("/summary")
def get_summary(
    start_date: Optional[date] = Query(None, description="YYYY-MM-DD (Premium only)"),
    end_date: Optional[date] = Query(None, description="YYYY-MM-DD (Premium only)"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Returns income/expense summary.
    - USER: always current month.
    - PREMIUM/ADMIN: respects custom date range.
    """
    data = analytics_service.get_summary(
        db=db,
        user_id=current_user.user_id,
        start_date=start_date,
        end_date=end_date,
        role=current_user.role
    )
    return data


# ==========================================
# GOALS PROGRESS
# (All roles)
# ==========================================

@router.get("/goals-progress")
def goals_progress(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Returns progress for all user's financial goals.
    Available to all roles.
    """
    return analytics_service.get_goals_progress(
        db=db,
        user_id=current_user.user_id
    )


# ==========================================
# MONTHLY TREND
# (PREMIUM + ADMIN only)
# ==========================================

@router.get("/monthly-trend")
def monthly_trend(
    months: int = Query(12, ge=3, le=24, description="Number of months to look back"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_premium_or_admin)
):
    """
    Returns income vs expense for each of the last N months.
    Premium/Admin only — returns 403 for basic USER.
    """
    return analytics_service.get_monthly_trend(
        db=db,
        user_id=current_user.user_id,
        months=months
    )


# ==========================================
# CATEGORY SPENDING OVER TIME
# (PREMIUM + ADMIN only)
# ==========================================

@router.get("/category-over-time")
def category_over_time(
    months: int = Query(6, ge=2, le=12, description="Number of months to look back"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_premium_or_admin)
):
    """
    Returns spending per category per month for the last N months.
    Premium/Admin only.
    """
    return analytics_service.get_category_over_time(
        db=db,
        user_id=current_user.user_id,
        months=months
    )


# ==========================================
# MONTH COMPARISON
# (PREMIUM + ADMIN only)
# ==========================================

@router.get("/comparison")
def comparison(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_premium_or_admin)
):
    """
    Returns current month vs previous month comparison of income/expenses.
    Premium/Admin only.
    """
    return analytics_service.get_comparison(
        db=db,
        user_id=current_user.user_id
    )


# ==========================================
# EXPORT ANALYTICS — EXCEL
# (PREMIUM + ADMIN only)
# ==========================================

@router.get("/export/excel")
def export_analytics_excel(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_premium_or_admin)
):
    """
    Exports current analytics data as an Excel file.
    Uses selected date range if provided; otherwise uses current month.
    Premium/Admin only.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    summary = analytics_service.get_summary(db, current_user.user_id, start_date, end_date, current_user.role)
    categories = analytics_service.get_spending_by_category(db, current_user.user_id, start_date, end_date, current_user.role)
    trend = analytics_service.get_monthly_trend(db, current_user.user_id, months=6)
    goals = analytics_service.get_goals_progress(db, current_user.user_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Analytics Report"
    ws.sheet_view.showGridLines = False

    TEAL = "1A6B5A"
    thin = Side(style="thin", color="C0C0C0")

    def border():
        return Border(top=thin, bottom=thin, left=thin, right=thin)

    def hdr_style(bg="1A6B5A"):
        return {
            "font": Font(name="Calibri", bold=True, size=11, color="FFFFFF"),
            "fill": PatternFill("solid", fgColor=bg),
            "alignment": Alignment(horizontal="center", vertical="center"),
            "border": border()
        }

    def write_cell(ws, row, col, value, **styles):
        cell = ws.cell(row=row, column=col, value=value)
        if "font" in styles: cell.font = styles["font"]
        if "fill" in styles: cell.fill = styles["fill"]
        if "alignment" in styles: cell.alignment = styles["alignment"]
        if "border" in styles: cell.border = styles["border"]
        return cell

    row = 1
    # Title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    tc = ws.cell(row=row, column=1, value="  BudgetBuddy Analytics Report")
    tc.font = Font(name="Calibri", bold=True, size=15, color="FFFFFF")
    tc.fill = PatternFill("solid", fgColor=TEAL)
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 30
    row += 1

    # Meta
    for label, val in [
        ("User", current_user.name),
        ("Period", summary["period_label"]),
        ("Generated", str(date.today()))
    ]:
        ws.cell(row=row, column=1, value=label + ":").font = Font(bold=True, color=TEAL)
        ws.cell(row=row, column=2, value=val)
        row += 1

    row += 1

    # Summary Section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    tc = ws.cell(row=row, column=1, value="  FINANCIAL SUMMARY")
    tc.font = Font(bold=True, size=12, color="FFFFFF")
    tc.fill = PatternFill("solid", fgColor=TEAL)
    tc.alignment = Alignment(horizontal="left", vertical="center")
    row += 1

    for metric, value, color in [
        ("Total Income", f"₹{summary['total_income']:,.2f}", "1A7A3C"),
        ("Total Expenses", f"₹{summary['total_expenses']:,.2f}", "C0392B"),
        ("Net Balance", f"₹{summary['net_balance']:,.2f}", TEAL if summary['net_balance'] >= 0 else "C0392B"),
        ("Savings Rate", f"{summary['savings_rate']}%", "1A7A3C")
    ]:
        lc = ws.cell(row=row, column=1, value=metric)
        lc.font = Font(bold=True, name="Calibri")
        lc.border = border()
        vc = ws.cell(row=row, column=2, value=value)
        vc.font = Font(bold=True, color=color, name="Calibri")
        vc.border = border()
        row += 1

    row += 1

    # Category Section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    hc = ws.cell(row=row, column=1, value="  SPENDING BY CATEGORY")
    hc.font = Font(bold=True, size=12, color="FFFFFF")
    hc.fill = PatternFill("solid", fgColor="2D9E7F")
    hc.alignment = Alignment(horizontal="left")
    row += 1

    for h, col in zip(["Category", "Amount", "% of Total"], [1, 2, 3]):
        write_cell(ws, row, col, h, **hdr_style("2D9E7F"))
    row += 1

    for cat in categories:
        ws.cell(row=row, column=1, value=cat["category"]).border = border()
        ws.cell(row=row, column=2, value=f"₹{cat['amount']:,.2f}").border = border()
        ws.cell(row=row, column=3, value=f"{cat['percentage']}%").border = border()
        row += 1

    row += 1

    # Goals Section
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    gc = ws.cell(row=row, column=1, value="  SAVINGS GOALS PROGRESS")
    gc.font = Font(bold=True, size=12, color="FFFFFF")
    gc.fill = PatternFill("solid", fgColor="6366F1")
    gc.alignment = Alignment(horizontal="left")
    row += 1

    for h, col in zip(["Goal", "Target", "Saved", "Progress"], [1, 2, 3, 4]):
        write_cell(ws, row, col, h, **hdr_style("6366F1"))
    row += 1

    for g in goals:
        ws.cell(row=row, column=1, value=g["goal_name"]).border = border()
        ws.cell(row=row, column=2, value=f"₹{g['target_amount']:,.2f}").border = border()
        ws.cell(row=row, column=3, value=f"₹{g['current_amount']:,.2f}").border = border()
        ws.cell(row=row, column=4, value=f"{g['progress_percentage']}%").border = border()
        row += 1

    # Column widths
    for col, width in zip("ABCD", [30, 22, 22, 18]):
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"BudgetBuddy_Analytics_{date.today()}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ==========================================
# EXPORT ANALYTICS — PDF
# (PREMIUM + ADMIN only)
# ==========================================

@router.get("/export/pdf")
def export_analytics_pdf(
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_premium_or_admin)
):
    """
    Exports analytics data as a PDF report.
    Premium/Admin only.
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.colors import HexColor, white, black
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="PDF generation library (reportlab) is not installed. Please install it: pip install reportlab"
        )

    summary = analytics_service.get_summary(db, current_user.user_id, start_date, end_date, current_user.role)
    categories = analytics_service.get_spending_by_category(db, current_user.user_id, start_date, end_date, current_user.role)
    goals = analytics_service.get_goals_progress(db, current_user.user_id)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2*cm, leftMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm
    )

    styles = getSampleStyleSheet()
    TEAL = HexColor("#1A6B5A")
    PURPLE = HexColor("#6366F1")

    title_style = ParagraphStyle(
        "Title", parent=styles["Title"],
        textColor=TEAL, fontSize=20, spaceAfter=6
    )
    heading_style = ParagraphStyle(
        "Heading", parent=styles["Heading2"],
        textColor=TEAL, fontSize=13, spaceAfter=4
    )
    normal = styles["Normal"]

    elements = []
    elements.append(Paragraph("💰 BudgetBuddy Analytics Report", title_style))
    elements.append(Paragraph(f"User: {current_user.name}  |  Period: {summary['period_label']}  |  Generated: {date.today()}", normal))
    elements.append(Spacer(1, 0.4*cm))

    # Summary Table
    elements.append(Paragraph("Financial Summary", heading_style))
    summary_data = [
        ["Metric", "Value"],
        ["Total Income", f"₹{summary['total_income']:,.2f}"],
        ["Total Expenses", f"₹{summary['total_expenses']:,.2f}"],
        ["Net Balance", f"₹{summary['net_balance']:,.2f}"],
        ["Savings Rate", f"{summary['savings_rate']}%"],
    ]
    st = Table(summary_data, colWidths=[10*cm, 7*cm])
    st.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), TEAL),
        ("TEXTCOLOR", (0, 0), (-1, 0), white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#F0FAF7"), white]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(st)
    elements.append(Spacer(1, 0.5*cm))

    # Category Breakdown Table
    elements.append(Paragraph("Spending by Category", heading_style))
    if categories:
        cat_data = [["Category", "Amount", "% of Total"]]
        for c in categories:
            cat_data.append([c["category"], f"₹{c['amount']:,.2f}", f"{c['percentage']}%"])
        ct = Table(cat_data, colWidths=[9*cm, 5*cm, 3*cm])
        ct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), PURPLE),
            ("TEXTCOLOR", (0, 0), (-1, 0), white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#F5F4FF"), white]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(ct)
    else:
        elements.append(Paragraph("No expense records for this period.", normal))
    elements.append(Spacer(1, 0.5*cm))

    # Goals Table
    elements.append(Paragraph("Savings Goals Progress", heading_style))
    if goals:
        goal_data = [["Goal", "Target", "Saved", "Progress"]]
        for g in goals:
            goal_data.append([
                g["goal_name"],
                f"₹{g['target_amount']:,.2f}",
                f"₹{g['current_amount']:,.2f}",
                f"{g['progress_percentage']}%"
            ])
        gt = Table(goal_data, colWidths=[8*cm, 4*cm, 4*cm, 3*cm])
        gt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), HexColor("#2D9E7F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#F0FAF5"), white]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(gt)
    else:
        elements.append(Paragraph("No savings goals found.", normal))

    elements.append(Spacer(1, 1*cm))
    elements.append(Paragraph(
        f"<i>BudgetBuddy Platform  •  Premium Analytics Report  •  {datetime.now().strftime('%d %b %Y %H:%M')}</i>",
        ParagraphStyle("footer", parent=normal, textColor=HexColor("#888888"), fontSize=8, alignment=1)
    ))

    doc.build(elements)
    buffer.seek(0)

    filename = f"BudgetBuddy_Analytics_{date.today()}.pdf"
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
