import io
from datetime import date
from fastapi import APIRouter, Depends, Query, Response
from sqlalchemy import extract, func
from sqlalchemy.orm import Session

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models.account import Account
from app.models.budget import Budget
from app.models.category import Category
from app.models.expense import Expense
from app.models.financial_goal import FinancialGoal
from app.models.income import Income
from app.models.user import User
from app.routers.auth import get_current_user


router = APIRouter(
    prefix="/reports",
    tags=["Reports"]
)


@router.get("/summary")
def get_report_summary(
    month: str | None = Query(None, description="Month in YYYY-MM format"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    user_id = current_user.user_id

    income_query = db.query(Income).filter(Income.user_id == user_id)
    expense_query = db.query(Expense).filter(Expense.user_id == user_id)

    filter_year = None
    filter_month = None

    if month:
        try:
            parts = month.split("-")
            filter_year = int(parts[0])
            filter_month = int(parts[1])

            income_query = income_query.filter(
                extract("year", Income.income_date) == filter_year,
                extract("month", Income.income_date) == filter_month
            )
            expense_query = expense_query.filter(
                extract("year", Expense.expense_date) == filter_year,
                extract("month", Expense.expense_date) == filter_month
            )
        except Exception:
            pass

    incomes = income_query.order_by(Income.income_date.desc()).all()
    expenses = expense_query.order_by(Expense.expense_date.desc()).all()
    accounts = db.query(Account).filter(Account.user_id == user_id).all()
    categories = db.query(Category).filter(Category.user_id == user_id).all()
    cat_map = {c.category_id: c.name for c in categories}

    total_income = sum(float(i.amount) for i in incomes)
    total_expenses = sum(float(e.amount) for e in expenses)
    savings = total_income - total_expenses

    # Category breakdown
    cat_totals = {}
    for exp in expenses:
        name = cat_map.get(exp.category_id, f"Category {exp.category_id}")
        cat_totals[name] = cat_totals.get(name, 0.0) + float(exp.amount)

    category_summary = [
        {
            "category_name": cat,
            "amount": amt,
            "percentage": round((amt / total_expenses * 100), 1) if total_expenses > 0 else 0.0
        }
        for cat, amt in cat_totals.items()
    ]
    category_summary.sort(key=lambda x: x["amount"], reverse=True)

    # 5. Combined Complete Transaction History
    combined_transactions = []
    for inc in incomes:
        combined_transactions.append({
            "transaction_id": f"inc-{inc.income_id}",
            "date": str(inc.income_date),
            "type": "Income",
            "category_or_source": inc.source or "Income",
            "bank_or_account": inc.bank_name or "-",
            "description": inc.description or "-",
            "amount": float(inc.amount),
            "formatted_amount": f"+ ₹{float(inc.amount):,.2f}",
            "raw_date": inc.income_date,
            "raw_id": inc.income_id
        })
    for exp in expenses:
        cname = cat_map.get(exp.category_id, f"Category {exp.category_id}")
        combined_transactions.append({
            "transaction_id": f"exp-{exp.expense_id}",
            "date": str(exp.expense_date),
            "type": "Expense",
            "category_or_source": cname,
            "bank_or_account": "-",
            "description": exp.description or "-",
            "amount": float(exp.amount),
            "formatted_amount": f"- ₹{float(exp.amount):,.2f}",
            "raw_date": exp.expense_date,
            "raw_id": exp.expense_id
        })

    # Sort newest transactions first (chronological descending)
    combined_transactions.sort(key=lambda t: (t["raw_date"], t["raw_id"]), reverse=True)

    # Budget info
    budgets = db.query(Budget).filter(Budget.user_id == user_id).all()
    selected_budget = None
    if filter_year and filter_month:
        for b in budgets:
            if b.month and b.month.year == filter_year and b.month.month == filter_month:
                selected_budget = b
                break

    budget_amt = float(selected_budget.amount) if selected_budget else 0.0

    return {
        "period": month if month else "All Time",
        "user": {
            "name": current_user.name,
            "email": current_user.email
        },
        "summary": {
            "total_income": total_income,
            "total_expenses": total_expenses,
            "net_savings": savings,
            "savings_rate": round((savings / total_income * 100), 1) if total_income > 0 else 0.0,
            "budget_amount": budget_amt,
            "budget_utilization": round((total_expenses / budget_amt * 100), 1) if budget_amt > 0 else 0.0
        },
        "category_summary": category_summary,
        "accounts": [
            {
                "account_name": a.account_name,
                "account_type": a.account_type,
                "balance": float(a.balance)
            }
            for a in accounts
        ],
        "incomes": [
            {
                "income_id": i.income_id,
                "amount": float(i.amount),
                "source": i.source,
                "bank_name": i.bank_name,
                "description": i.description,
                "income_date": str(i.income_date)
            }
            for i in incomes
        ],
        "expenses": [
            {
                "expense_id": e.expense_id,
                "amount": float(e.amount),
                "category_name": cat_map.get(e.category_id, f"Category {e.category_id}"),
                "description": e.description,
                "expense_date": str(e.expense_date)
            }
            for e in expenses
        ],
        "transactions": [
            {
                "id": t["transaction_id"],
                "date": t["date"],
                "type": t["type"],
                "categoryOrSource": t["category_or_source"],
                "bankOrAccount": t["bank_or_account"],
                "description": t["description"],
                "amount": t["amount"],
                "formatted_amount": t["formatted_amount"]
            }
            for t in combined_transactions
        ]
    }


@router.get("/export/excel")
@router.get("/export/csv")          # keep old URL so frontend needs zero changes
def export_excel(
    month: str | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Generate and stream a professional Excel (.xlsx) financial report.
    Keeps backward-compatible URL /export/csv so the frontend needs no changes.
    """
    user_id = current_user.user_id

    # ── 1. Fetch data ────────────────────────────────────────
    income_query = db.query(Income).filter(Income.user_id == user_id)
    expense_query = db.query(Expense).filter(Expense.user_id == user_id)

    if month:
        try:
            parts = month.split("-")
            f_year = int(parts[0])
            f_month = int(parts[1])
            income_query = income_query.filter(
                extract("year", Income.income_date) == f_year,
                extract("month", Income.income_date) == f_month
            )
            expense_query = expense_query.filter(
                extract("year", Expense.expense_date) == f_year,
                extract("month", Expense.expense_date) == f_month
            )
        except Exception:
            pass

    incomes = income_query.order_by(Income.income_date.desc(), Income.income_id.desc()).all()
    expenses = expense_query.order_by(Expense.expense_date.desc(), Expense.expense_id.desc()).all()
    categories = db.query(Category).filter(Category.user_id == user_id).all()
    cat_map = {c.category_id: c.name for c in categories}

    # Build combined transaction history (newest first)
    combined_transactions = []
    for inc in incomes:
        combined_transactions.append({
            "date": str(inc.income_date),
            "type": "Income",
            "category_or_source": inc.source or "Income",
            "bank_or_account": inc.bank_name or "-",
            "description": inc.description or "-",
            "amount_num": float(inc.amount),
            "amount_str": f"+ INR {float(inc.amount):,.2f}",
            "raw_date": inc.income_date,
            "raw_id": inc.income_id,
        })
    for exp in expenses:
        cname = cat_map.get(exp.category_id, f"Category {exp.category_id}")
        combined_transactions.append({
            "date": str(exp.expense_date),
            "type": "Expense",
            "category_or_source": cname,
            "bank_or_account": "-",
            "description": exp.description or "-",
            "amount_num": float(exp.amount),
            "amount_str": f"- INR {float(exp.amount):,.2f}",
            "raw_date": exp.expense_date,
            "raw_id": exp.expense_id,
        })
    combined_transactions.sort(key=lambda t: (t["raw_date"], t["raw_id"]), reverse=True)

    tot_inc = sum(float(i.amount) for i in incomes)
    tot_exp = sum(float(e.amount) for e in expenses)
    net_savings = tot_inc - tot_exp

    # ── 2. Style helpers ─────────────────────────────────────
    thin = Side(style="thin", color="C0C0C0")

    def cell_border():
        return Border(top=thin, bottom=thin, left=thin, right=thin)

    def apply_style(cell, font=None, fill=None, alignment=None, border=None):
        if font:       cell.font       = font
        if fill:       cell.fill       = fill
        if alignment:  cell.alignment  = alignment
        if border:     cell.border     = border

    # Shared colours
    DARK_TEAL   = "1A6B5A"
    MID_TEAL    = "2D9E7F"
    INCOME_BG   = "D6F5E3"
    EXPENSE_BG  = "FAE3E3"
    SUMMARY_BG  = "FFF9E6"

    def section_heading_style():
        return dict(
            font=Font(name="Calibri", bold=True, size=13, color="FFFFFF"),
            fill=PatternFill("solid", fgColor=DARK_TEAL),
            alignment=Alignment(horizontal="left", vertical="center"),
        )

    def col_header_style(bg_color):
        return dict(
            font=Font(name="Calibri", bold=True, size=10, color="FFFFFF"),
            fill=PatternFill("solid", fgColor=bg_color),
            alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=cell_border(),
        )

    def data_cell_style(bg_color, bold=False, color="1C1C1C", align="left"):
        return dict(
            font=Font(name="Calibri", bold=bold, size=10, color=color),
            fill=PatternFill("solid", fgColor=bg_color),
            alignment=Alignment(horizontal=align, vertical="center", wrap_text=True),
            border=cell_border(),
        )

    # ── 3. Build workbook ────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Report"
    ws.sheet_view.showGridLines = False

    NCOLS = 6
    period_label = month if month else "All Time"
    generated_date = str(date.today())

    row = 1

    # Title banner
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
    title_cell = ws.cell(row=row, column=1, value="   BudgetBuddy Financial Report")
    apply_style(
        title_cell,
        font=Font(name="Calibri", bold=True, size=16, color="FFFFFF"),
        fill=PatternFill("solid", fgColor=DARK_TEAL),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    ws.row_dimensions[row].height = 34
    row += 1

    # Meta rows
    meta_rows = [
        ("Account Holder", current_user.name),
        ("Email",          current_user.email),
        ("Generated Date", generated_date),
        ("Statement Period", period_label),
    ]
    for label, value in meta_rows:
        lc = ws.cell(row=row, column=1, value=label + ":")
        vc = ws.cell(row=row, column=2, value=value)
        apply_style(lc, font=Font(name="Calibri", bold=True, size=11, color=DARK_TEAL),
                    alignment=Alignment(horizontal="left", vertical="center"))
        apply_style(vc, font=Font(name="Calibri", size=11, color="1C1C1C"),
                    alignment=Alignment(horizontal="left", vertical="center"))
        ws.row_dimensions[row].height = 16
        row += 1

    row += 1  # blank gap

    # ─── Helper: write section heading ───────────────────────
    def write_heading(text):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
        cell = ws.cell(row=row, column=1, value=f"  {text}")
        apply_style(cell, **section_heading_style())
        ws.row_dimensions[row].height = 22
        row += 1

    # ─── Helper: write column header row ─────────────────────
    def write_col_headers(headers, bg_color):
        nonlocal row
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=c, value=h)
            apply_style(cell, **col_header_style(bg_color))
        ws.row_dimensions[row].height = 18
        row += 1

    # ─── Helper: write a data row ────────────────────────────
    def write_data_row(values, bg_color, amount_col=None,
                       amount_color="1A7A3C"):
        nonlocal row
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=val)
            if amount_col and c == amount_col:
                apply_style(cell,
                             font=Font(name="Calibri", bold=True, size=10, color=amount_color),
                             fill=PatternFill("solid", fgColor=bg_color),
                             alignment=Alignment(horizontal="right", vertical="center"),
                             border=cell_border())
            else:
                apply_style(cell, **data_cell_style(bg_color))
        ws.row_dimensions[row].height = 15
        row += 1

    # ═══════════════════════════════════════════════════════════
    #  SECTION A — INCOME RECORDS
    # ═══════════════════════════════════════════════════════════
    write_heading("INCOME RECORDS")
    write_col_headers(["Date", "Source", "Bank / Account", "Amount (INR)", "Description"], MID_TEAL)

    if incomes:
        for inc in incomes:
            write_data_row(
                [str(inc.income_date), inc.source or "-", inc.bank_name or "-",
                 f"+ INR {float(inc.amount):,.2f}", inc.description or "-"],
                INCOME_BG, amount_col=4, amount_color="1A7A3C"
            )
    else:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row=row, column=1, value="No income records for this period.")
        row += 1

    row += 1  # gap

    # ═══════════════════════════════════════════════════════════
    #  SECTION B — EXPENSE RECORDS
    # ═══════════════════════════════════════════════════════════
    write_heading("EXPENSE RECORDS")
    write_col_headers(["Date", "Category", "Amount (INR)", "Description"], "C0392B")

    if expenses:
        for exp in expenses:
            cname = cat_map.get(exp.category_id, f"Category {exp.category_id}")
            write_data_row(
                [str(exp.expense_date), cname,
                 f"- INR {float(exp.amount):,.2f}", exp.description or "-"],
                EXPENSE_BG, amount_col=3, amount_color="C0392B"
            )
    else:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row=row, column=1, value="No expense records for this period.")
        row += 1

    row += 1  # gap

    # ═══════════════════════════════════════════════════════════
    #  SECTION C — COMPLETE TRANSACTION HISTORY
    # ═══════════════════════════════════════════════════════════
    write_heading("COMPLETE TRANSACTION HISTORY")
    history_header_row = row
    write_col_headers(
        ["Date", "Transaction Type", "Category / Source",
         "Bank / Account", "Description", "Amount (INR)"],
        "3F51B5"
    )

    # Freeze pane just below the transaction header
    ws.freeze_panes = ws.cell(row=history_header_row + 1, column=1)

    if combined_transactions:
        for t in combined_transactions:
            is_income = t["type"] == "Income"
            bg = INCOME_BG if is_income else EXPENSE_BG
            amt_color = "1A7A3C" if is_income else "C0392B"
            write_data_row(
                [t["date"], t["type"], t["category_or_source"],
                 t["bank_or_account"], t["description"], t["amount_str"]],
                bg, amount_col=6, amount_color=amt_color
            )
    else:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
        ws.cell(row=row, column=1, value="No transactions available for this statement period.")
        row += 1

    row += 1  # gap

    # ═══════════════════════════════════════════════════════════
    #  SECTION D — FINANCIAL SUMMARY
    # ═══════════════════════════════════════════════════════════
    write_heading("FINANCIAL SUMMARY")
    write_col_headers(["Metric", "Amount (INR)"], "E67E22")

    summary_rows = [
        ("Total Income",   f"INR {tot_inc:,.2f}",     "1A7A3C"),
        ("Total Expenses", f"INR {tot_exp:,.2f}",     "C0392B"),
        ("Net Savings",    f"INR {net_savings:,.2f}", "1A6B5A" if net_savings >= 0 else "C0392B"),
    ]
    for label, value, color in summary_rows:
        lc = ws.cell(row=row, column=1, value=label)
        vc = ws.cell(row=row, column=2, value=value)
        apply_style(lc,
                     font=Font(name="Calibri", bold=True, size=11, color="1C1C1C"),
                     fill=PatternFill("solid", fgColor=SUMMARY_BG),
                     alignment=Alignment(horizontal="left", vertical="center"),
                     border=cell_border())
        apply_style(vc,
                     font=Font(name="Calibri", bold=True, size=11, color=color),
                     fill=PatternFill("solid", fgColor=SUMMARY_BG),
                     alignment=Alignment(horizontal="right", vertical="center"),
                     border=cell_border())
        ws.row_dimensions[row].height = 18
        row += 1

    # Footer
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
    footer = ws.cell(
        row=row, column=1,
        value="BudgetBuddy Platform  •  Automated Student Financial Planning Report  •  End of Statement"
    )
    footer.font      = Font(name="Calibri", italic=True, size=9, color="888888")
    footer.alignment = Alignment(horizontal="center")

    # ── Professional column widths (no ####### ever) ─────────
    col_widths = [
        ("A", 14),   # Date
        ("B", 22),   # Source / Type / Metric
        ("C", 26),   # Bank / Category / Source
        ("D", 26),   # Description / Bank
        ("E", 36),   # Description
        ("F", 22),   # Amount
    ]
    for col_letter, width in col_widths:
        ws.column_dimensions[col_letter].width = width

    # ── Stream to bytes ───────────────────────────────────────
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"BudgetBuddy_Report_{month if month else 'All_Time'}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
